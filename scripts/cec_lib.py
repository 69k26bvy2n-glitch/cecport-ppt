# -*- coding: utf-8 -*-
"""中电港 · 咨询报告信息图风 —— 原生可编辑 PPTX 组件库
着色规则：形状一律纯色（蓝阶 NAVY/BLUE/SKY + 橙 ORANGE）；
渐变只用于蓝→橙装饰细线（set_gradient），形状不做自身渐变、不与色阶叠加。
中文排版：段落开启避头尾（eaLnBrk）+ 标点悬挂（hangingPunct），run 标 zh-CN；
不写 presentation 级 kinsoku 表（PowerPoint 对自注字符表校验严格易触发修复，
内置中文规则已覆盖常用标点）。"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import os

# ---------------- tokens ----------------
BG      = RGBColor(0xF3, 0xF5, 0xF8)
CARD    = RGBColor(0xFF, 0xFF, 0xFF)
HAIR    = RGBColor(0xE2, 0xE8, 0xF0)
# 4 级蓝阶（2026-10-14 重定：整体提饱和、去发灰淡蓝；深→浅）
NAVY    = RGBColor(0x00, 0x57, 0xC8)   # hsl 214/100%/39% 标题·一级
BLUE    = RGBColor(0x14, 0x78, 0xE0)   # hsl 211/84%/48% 二级
SKY     = RGBColor(0x45, 0x99, 0xF0)   # hsl 210/85%/61% 三级
MIST    = RGBColor(0x6A, 0xAE, 0xF3)   # hsl 210/85%/68% 四级（仍饱和，不发灰）
TINT    = RGBColor(0xE7, 0xF0, 0xF8)
TINT2   = RGBColor(0xF2, 0xF7, 0xFC)
ORANGE  = RGBColor(0xE8, 0x6F, 0x0B)
ORANGE2 = RGBColor(0xC9, 0x5E, 0x09)
TINT_OR = RGBColor(0xFB, 0xEF, 0xE3)
OR_LINE = RGBColor(0xF3, 0xD3, 0xB4)
INK     = RGBColor(0x20, 0x29, 0x36)
SUB     = RGBColor(0x5C, 0x66, 0x75)
MUTE    = RGBColor(0x93, 0xA0, 0xAF)
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)

# 蓝色阶（纯色，并列同类元素由深到浅取色）
BLUE_SCALE = [NAVY, BLUE, SKY]

LATIN = "MiSans"
EA    = "MiSans"

W, H = 13.333, 7.5
ML, MR = 0.55, 0.55
CW = W - ML - MR

_HERE = os.path.dirname(os.path.abspath(__file__))
_ASSETS = None
for _cand in (os.path.join(_HERE, "assets"),
              os.path.join(_HERE, "..", "assets"),
              os.path.join(os.getcwd(), "assets")):
    if os.path.isdir(_cand):
        _ASSETS = os.path.abspath(_cand); break
if _ASSETS is None:
    _ASSETS = os.path.join(_HERE, "assets")
LOGO   = os.path.join(_ASSETS, "logo.png")
FOOTER = os.path.join(_ASSETS, "footer.png")
ICONS  = os.path.join(_ASSETS, "icons2/")
ICONS_W = os.path.join(_ASSETS, "icons2w/")
FBOX = dict(x=511628/914400, y=6543487/914400, w=11079491/914400, h=149095/914400)

# ---------------- low level ----------------
def set_font(run, size=None, bold=None, color=None, latin=LATIN, ea=EA):
    f = run.font
    if size is not None: f.size = Pt(size)
    if bold is not None: f.bold = bold
    if color is not None: f.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    rPr.set("lang", "zh-CN")
    rPr.set("altLang", "en-US")
    for tag, face in (("a:latin", latin), ("a:ea", ea), ("a:cs", latin)):
        e = rPr.find(qn(tag))
        if e is None:
            e = rPr.makeelement(qn(tag), {}); rPr.append(e)
        e.set("typeface", face)

def soft_shadow(shape, blur=0.10, dist=0.035, alpha=13, color="1B2A41"):
    spPr = shape._element.spPr
    for old in spPr.findall(qn("a:effectLst")):
        spPr.remove(old)
    eff = spPr.makeelement(qn("a:effectLst"), {})
    sh = eff.makeelement(qn("a:outerShdw"), {
        "blurRad": str(int(blur * 914400)), "dist": str(int(dist * 914400)),
        "dir": "5400000", "rotWithShape": "0"})
    c = sh.makeelement(qn("a:srgbClr"), {"val": color})
    a = c.makeelement(qn("a:alpha"), {"val": str(alpha * 1000)})
    c.append(a); sh.append(c); eff.append(sh)
    ln = spPr.find(qn("a:ln"))
    if ln is not None: ln.addnext(eff)
    else: spPr.append(eff)

def add_rect(slide, x, y, w, h, fill=None, line=None, line_w=0.75,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.055, shadow=False):
    sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None: sp.fill.background()
    else: sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None: sp.line.fill.background()
    else: sp.line.color.rgb = line; sp.line.width = Pt(line_w)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try: sp.adjustments[0] = radius
        except Exception: pass
    if shadow: soft_shadow(sp)
    sp.shadow.inherit = False
    return sp

# 圆角统一语言（绝对英寸）。坑：ROUNDED_RECTANGLE 的 adjustments[0] 是
# 「圆角半径 / 短边」的比例而非英寸——同写 0.10、高度不同的框实际弧度会差一倍，
# 必须用 rr() 换算。三档：外卡/大容器 0.10，内盒/记录卡 0.08，图标方块自定 0.12。
R_CARD = 0.10
R_BOX  = 0.08
def rr(w, h, r=R_BOX):
    return r / min(w, h)

def card_shadow(shape):
    """整页只有外层卡片带投影：轻柔、短距离、低透明，层级语言统一。
    卡内元素（药丸/标签/小盒）一律不加阴影，靠填色描边分层。"""
    soft_shadow(shape, blur=0.09, dist=0.022, alpha=9)

def add_text(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, wrap=True):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = p.get("align", PP_ALIGN.LEFT)
        if p.get("space_after") is not None: para.space_after = Pt(p["space_after"])
        if p.get("line_spacing"): para.line_spacing = p["line_spacing"]
        # 显式开启中文避头尾：后置标点不进行首，行尾标点允许悬挂溢出
        pPr = para._p.get_or_add_pPr()
        pPr.set("eaLnBrk", "1")
        pPr.set("latinLnBrk", "0")
        pPr.set("hangingPunct", "1")
        for text, o in p["runs"]:
            r = para.add_run(); r.text = text
            set_font(r, o.get("size", 11), o.get("bold", False), o.get("color", INK))
    return tb

def R(text, size=11, bold=False, color=INK):
    return (text, {"size": size, "bold": bold, "color": color})

def set_gradient(shape, c1, c2, angle=0):
    """仅用于装饰细线（蓝→橙），不要给色块/形状用"""
    spPr = shape._element.spPr
    for tag in ("a:noFill", "a:solidFill", "a:gradFill", "a:blipFill",
                "a:pattFill", "a:grpFill"):
        for e in spPr.findall(qn(tag)): spPr.remove(e)
    grad = spPr.makeelement(qn("a:gradFill"), {})
    lst = grad.makeelement(qn("a:gsLst"), {})
    for pos, col in ((0, c1), (100000, c2)):
        gs = grad.makeelement(qn("a:gs"), {"pos": str(pos)})
        cl = gs.makeelement(qn("a:srgbClr"), {"val": str(col)})
        gs.append(cl); lst.append(gs)
    grad.append(lst)
    grad.append(grad.makeelement(qn("a:lin"), {"ang": str(angle * 60000), "scaled": "1"}))
    ln = spPr.find(qn("a:ln"))
    if ln is not None: ln.addprevious(grad)
    else:
        eff = spPr.find(qn("a:effectLst"))
        (eff.addprevious(grad) if eff is not None else spPr.append(grad))

def fill_alpha(shape, pct):
    c = shape._element.spPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
    if c is not None:
        for a in c.findall(qn("a:alpha")): c.remove(a)
        c.append(c.makeelement(qn("a:alpha"), {"val": str(int(pct * 1000))}))

def add_line(slide, x, y, w, color=HAIR, weight=0.75):
    ln = slide.shapes.add_connector(1, Inches(x), Inches(y), Inches(x + w), Inches(y))
    ln.line.color.rgb = color; ln.line.width = Pt(weight)
    ln.shadow.inherit = False
    return ln

def add_line_xy(slide, x1, y1, x2, y2, color=MIST, weight=1.0):
    ln = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    ln.line.color.rgb = color; ln.line.width = Pt(weight)
    ln.shadow.inherit = False
    return ln

def _ln(shape):
    return shape.line._get_or_add_ln()

def set_dash(shape, val="dash"):
    """流程虚线：a:prstDash（dash/dashDot 等）。"""
    el = _ln(shape)
    el.append(el.makeelement(qn("a:prstDash"), {"val": val}))

def set_arrow(shape, head=False, tail=False):
    """三角箭头。坑：a:headEnd 装饰画线起点、a:tailEnd 才在画线终点；
    add_line_xy(x1,y1,x2,y2) 的终点是 (x2,y2)，要箭头指到 (x2,y2) 用 tail=True。
    双向流 head+tail 同开。"""
    el = _ln(shape)
    if head:
        el.append(el.makeelement(qn("a:headEnd"),
            {"type": "triangle", "w": "med", "len": "med"}))
    if tail:
        el.append(el.makeelement(qn("a:tailEnd"),
            {"type": "triangle", "w": "med", "len": "med"}))

def add_icon(slide, name, x, y, d, white=False):
    return slide.shapes.add_picture((ICONS_W if white else ICONS) + name + ".png",
                                    Inches(x), Inches(y), Inches(d), Inches(d))

def icon_chip(slide, name, cx, cy, d=0.52, bg=TINT):
    x, y = cx - d / 2, cy - d / 2
    add_rect(slide, x, y, d, d, fill=bg, radius=0.28)
    pad = d * 0.20
    add_icon(slide, name, x + pad, y + pad, d - 2 * pad)
    return x, y

# ---------------- 流程图：泳道与穿线标签 ----------------
def dashed_box(slide, x, y, w, h, fill=TINT2, line=SKY, weight=1.1):
    """泳道虚线容器：浅色实底 + 同色系 1.1pt 虚线描边，外卡圆角档。"""
    sp = add_rect(slide, x, y, w, h, fill=fill, line=line, line_w=weight,
                  radius=rr(w, h, R_CARD))
    set_dash(sp, "dash")
    return sp

def lane_title(slide, cx, y, w, text, fill, h=0.36):
    """泳道标题药丸：实底白字、无阴影（卡内元素不投影）。"""
    add_rect(slide, cx - w / 2, y, w, h, fill=fill, radius=0.5)
    add_text(slide, cx - w / 2, y - 0.005, w, h,
             [{"runs": [R(text, 11, True, WHITE)], "align": PP_ALIGN.CENTER}],
             anchor=MSO_ANCHOR.MIDDLE)

VL_W = 0.22   # 竖排穿线标签框宽（全角字 8pt≈0.11in + 两侧小余量）

def vlabel(slide, cx, cy, text, color=BLUE, size=8):
    """竖排穿线标签（列间隙放不下横排胶囊时用）：
    PPT 原生东亚竖排 a:bodyPr vert=eaVert（字形正立、自上而下成列，绝不旋转文本框）；
    文字框本身白色实填、直角、无描边，居中压在贯通连线上断线，文字色随线色。
    实线可整条贯通由白框压线；虚线必须画两段、端点切入白框 0.02in（否则框边外露虚线碎头）。"""
    n = len(text)
    hh = n * size / 72 * 1.22      # 列高 = n 个字
    tb = add_text(slide, cx - VL_W / 2, cy - hh / 2, VL_W, hh,
                  [{"runs": [R(text, size, True, color)],
                    "align": PP_ALIGN.CENTER}],
                  anchor=MSO_ANCHOR.MIDDLE, wrap=False)
    bp = tb.text_frame._txBody.find(qn("a:bodyPr"))
    bp.set("vert", "eaVert")
    # bodyPr 默认 lIns/tIns=0.1in，eaVert 下会把竖排列顶偏约 0.1in，必须清零
    for ins in ("lIns", "tIns", "rIns", "bIns"):
        bp.set(ins, "0")
    tb.fill.solid(); tb.fill.fore_color.rgb = CARD
    tb.line.fill.background()
    return tb

def line_pill(slide, cx, cy, w, text, size, color=BLUE, h=0.20, rot=0.0):
    """横排穿线标签：白色圆角胶囊 + 与线同色 1pt 描边，文字色随线色。
    实线可整条贯通由胶囊压线；虚线画两段、端点切入胶囊 0.02in。
    rot：斜线标签沿走向顺时针旋转角（度），胶囊与文字框同角、绕共同中心旋转；
    旋转后线从短边穿出，断口按 t=w/2-0.02 沿射线收。"""
    pill = add_rect(slide, cx - w / 2, cy - h / 2, w, h,
                    fill=CARD, line=color, line_w=1.0, radius=0.5)
    pill.rotation = rot
    tw = w - 0.153
    tb = add_text(slide, cx - tw / 2, cy - 0.074, tw, 0.148,
                  [{"runs": [R(text, size, True, color)], "align": PP_ALIGN.CENTER}],
                  anchor=MSO_ANCHOR.MIDDLE)
    tb.rotation = rot
    return pill

# ---------------- composite ----------------
def new_deck():
    prs = Presentation()
    prs.slide_width = Inches(W); prs.slide_height = Inches(H)
    # 自定义尺寸下删除模板残留的 4x3 类型标记，避免尺寸与类型矛盾
    sldsz = prs.part._element.find(qn("p:sldSz"))
    if sldsz is not None:
        del sldsz.attrib["type"]
    return prs

def new_page(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(slide, 0, 0, W, H, fill=BG, shape=MSO_SHAPE.RECTANGLE)
    o1 = add_rect(slide, -1.25, 6.05, 2.9, 2.9, fill=TINT, shape=MSO_SHAPE.OVAL)
    fill_alpha(o1, 35)
    o2 = add_rect(slide, 12.15, 6.55, 2.2, 2.2, fill=TINT_OR, shape=MSO_SHAPE.OVAL)
    fill_alpha(o2, 38)
    return slide

def header(slide, title, kicker=None):
    """页眉：单根蓝→橙渐变细竖线（0.07×0.72，圆头），纵向贯穿主标题与副标题。
    2026-10-18 用户定稿；禁止再退回"实心橙竖条+渐变横线"旧方案。"""
    bar = add_rect(slide, ML + 0.02, 0.40, 0.07, 0.72, radius=0.5)
    set_gradient(bar, BLUE, ORANGE, angle=90)
    add_text(slide, ML + 0.26, 0.36, 9.4, 0.52,
             [{"runs": [R(title, 23, True, NAVY)]}], anchor=MSO_ANCHOR.MIDDLE)
    if kicker:
        add_text(slide, ML + 0.26, 0.885, 9.4, 0.28,
                 [{"runs": [R(kicker, 10.5, False, MUTE)]}],
                 anchor=MSO_ANCHOR.MIDDLE)
    slide.shapes.add_picture(LOGO, Inches(W - MR - 1.92), Inches(0.40),
                             height=Inches(0.40))

def footer(slide, page=None):
    slide.shapes.add_picture(FOOTER, Inches(FBOX["x"]), Inches(FBOX["y"]),
                             width=Inches(FBOX["w"]), height=Inches(FBOX["h"]))
    if page is not None:
        add_text(slide, W - MR - 1.1, FBOX["y"] - 0.27, 1.1, 0.22,
                 [{"runs": [R(f"{page:02d} / 25", 8, False, MUTE)],
                   "align": PP_ALIGN.RIGHT}], anchor=MSO_ANCHOR.BOTTOM)

def card(slide, x, y, w, h, r=R_CARD):
    """外层卡片：白 + HAIR 边 + 绝对英寸圆角（默认 R_CARD=0.10）+ card_shadow 轻投影。
    卡内元素不要再加阴影。旧 radius=比例 关键字已停用，需要其他半径传 r=英寸值。"""
    sp = add_rect(slide, x, y, w, h, fill=CARD, line=HAIR, line_w=0.75,
                  radius=rr(w, h, r))
    card_shadow(sp)
    return sp

def rows_y(y0, y1, n, gap=0.12, min_h=0.68):
    """在 [y0, y1] 内等分 n 行，返回 [(y, h), ...]。
    用于卡内纵向列表（KPI/指标/记录行）：先算行高再摆，避免内容压出卡底。
    行高不足 min_h 时抛错——按 page-patterns §4.5 缩字号/减间距/拆页，不许硬塞。"""
    h = (y1 - y0 - (n - 1) * gap) / n
    if h < min_h:
        raise ValueError(
            f"行高 {h:.2f}in < {min_h}in（n={n}, 可用高 {y1 - y0:.2f}in）："
            "缩小字号、把 gap 降到 0.08，或拆页/改两栏")
    return [(y0 + i * (h + gap), h) for i in range(n)]

def card_title(slide, x, y, w, text, size=13.5, marker=ORANGE, icon=None):
    if icon:
        add_icon(slide, icon, x, y + 0.06, 0.26)
        tx = x + 0.38
    else:
        add_rect(slide, x, y + 0.10, 0.10, 0.10, fill=marker, radius=0.22)
        tx = x + 0.20
    add_text(slide, tx, y, w - (tx - x), 0.34,
             [{"runs": [R(text, size, True, NAVY)]}],
             anchor=MSO_ANCHOR.MIDDLE)

def icon_title(slide, x, y, w, text, icon, isize=0.26, size=13):
    add_icon(slide, icon, x, y + (0.34 - isize) / 2, isize)
    add_text(slide, x + isize + 0.10, y, w - isize - 0.1, 0.34,
             [{"runs": [R(text, size, True, NAVY)]}],
             anchor=MSO_ANCHOR.MIDDLE)

def badge(slide, x, y, d, num, fill=NAVY, tsize=15):
    add_rect(slide, x, y, d, d, fill=fill, shape=MSO_SHAPE.OVAL)
    add_text(slide, x, y - 0.01, d, d,
             [{"runs": [R(num, tsize, True, WHITE)], "align": PP_ALIGN.CENTER}],
             anchor=MSO_ANCHOR.MIDDLE)

def chip(slide, x, y, d, glyph, fill=BLUE, gsize=12):
    add_rect(slide, x, y, d, d, fill=fill, radius=0.22)
    add_text(slide, x, y - 0.01, d, d,
             [{"runs": [R(glyph, gsize, True, WHITE)], "align": PP_ALIGN.CENTER}],
             anchor=MSO_ANCHOR.MIDDLE)

def tag(slide, x, y, text, fill=TINT, color=BLUE, w=None, size=9):
    tw = w if w else (0.18 + len(text) * size * 0.013)
    add_rect(slide, x, y, tw, 0.26, fill=fill, radius=0.5)
    add_text(slide, x, y - 0.005, tw, 0.26,
             [{"runs": [R(text, size, True, color)], "align": PP_ALIGN.CENTER}],
             anchor=MSO_ANCHOR.MIDDLE)
    return tw
